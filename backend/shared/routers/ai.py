from __future__ import annotations

import json
from datetime import datetime
from typing import Annotated

from fastapi import APIRouter, Depends, Request
from fastapi.responses import JSONResponse
from sqlalchemy.orm import Session

from shared.database import get_db
from shared.groq_client import call_groq_api
from shared.models import AIChatMessage, AIChatSession, Attendance, Feedback, Mark, User
from shared.security import get_current_user, get_token_payload

router = APIRouter(tags=["ai"])

try:
    import docx
except ImportError:
    docx = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

def _fail(message: str, status: int = 400) -> JSONResponse:
    return JSONResponse({"success": False, "message": message}, status_code=status)


def _extract_file_text(file) -> tuple[str, str | None]:
    filename = (file.filename or "").lower()
    file_text = ""
    error = None

    try:
        if filename.endswith(".pdf"):
            try:
                import PyPDF2

                pdf_reader = PyPDF2.PdfReader(file.file)
                for page in pdf_reader.pages:
                    text = page.extract_text()
                    if text:
                        file_text += text + "\n"
            except ImportError:
                error = "Server missing PyPDF2 library. Please install it."
        elif filename.endswith((".docx", ".doc")):
            if not docx:
                error = "Server missing python-docx library."
            else:
                try:
                    doc = docx.Document(file.file)
                    for para in doc.paragraphs:
                        file_text += para.text + "\n"
                except Exception as e:
                    file_text = f"[Error reading DOCX: {str(e)}]"
        elif filename.endswith(".pptx"):
            if not Presentation:
                error = "Server missing python-pptx library."
            else:
                try:
                    prs = Presentation(file.file)
                    for slide in prs.slides:
                        for shape in slide.shapes:
                            if hasattr(shape, "text"):
                                file_text += shape.text + "\n"
                except Exception as e:
                    file_text = f"[Error reading PPTX: {str(e)}]"
        elif filename.endswith((".txt", ".md", ".py", ".js", ".html", ".css", ".json")):
            file_text = file.file.read().decode("utf-8", errors="ignore")
        else:
            file_text = f"[Uploaded file: {file.filename} - Type not supported for automatic reading]"
    except Exception as e:
        print(f"File read error: {e}")
        file_text = f"[Error reading file: {str(e)}]"

    return file_text, error


@router.post("/chat")
async def chat(
    request: Request,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = ""
    sid = None
    file_text = ""
    uploaded_filename = None

    content_type = request.headers.get("content-type", "")
    if "application/json" in content_type:
        data = await request.json()
        query = data.get("question") or data.get("q") or ""
        sid = data.get("session_id")
    else:
        form = await request.form()
        query = form.get("question") or form.get("q") or ""
        sid = form.get("session_id")
        upload = form.get("file")
        if upload and hasattr(upload, "filename") and upload.filename:
            uploaded_filename = upload.filename
            file_text, err = _extract_file_text(upload)
            if err:
                return _fail(err, 500)

    if not query and not file_text:
        return _fail("Question or file required", 400)

    user_id = current_user.id
    session = None
    if sid:
        session = db.query(AIChatSession).filter_by(id=sid, user_id=user_id).first()
        if not session:
            return _fail("Session not found or access denied", 404)
        session.updated_at = datetime.utcnow()
    else:
        title = query[:50] + "..." if query else "File Analysis"
        session = AIChatSession(user_id=user_id, title=title)
        db.add(session)
        db.commit()

    user_msg_text = query
    if file_text:
        user_msg_text += f"\n[Attached: {uploaded_filename or 'File'}]"
    db.add(AIChatMessage(session_id=session.id, role="user", text=user_msg_text))

    final_prompt = query
    if file_text:
        final_prompt += (
            f"\n\n--- CONTEXT FROM UPLOADED FILE ({uploaded_filename}) ---\n"
            f"{file_text[:20000]}\n--- END CONTEXT ---\n"
            "(Note: Text has been extracted from the file. Answer based on this context if relevant.)"
        )

    answer, error_msg = call_groq_api(final_prompt)
    if answer:
        db.add(AIChatMessage(session_id=session.id, role="ai", text=answer))
        db.commit()
        return {"success": True, "answer": answer, "session_id": session.id}

    return JSONResponse(
        {"success": False, "message": error_msg or "Failed to get response from AI model."},
        status_code=500,
    )


@router.get("/ai/sessions")
def get_ai_sessions(
    payload: Annotated[dict, Depends(get_token_payload)],
    db: Session = Depends(get_db),
):
    uid = int(payload.get("sub"))
    sessions = (
        db.query(AIChatSession)
        .filter_by(user_id=uid)
        .order_by(AIChatSession.updated_at.desc())
        .all()
    )
    out = [{"id": s.id, "title": s.title, "updated_at": s.updated_at.isoformat()} for s in sessions]
    return {"success": True, "sessions": out}


@router.get("/ai/session/{session_id}")
def get_ai_session(
    session_id: str,
    payload: Annotated[dict, Depends(get_token_payload)],
    db: Session = Depends(get_db),
):
    uid = int(payload.get("sub"))
    session = db.query(AIChatSession).filter_by(id=session_id, user_id=uid).first()
    if not session:
        return _fail("Session not found", 404)

    messages = (
        db.query(AIChatMessage)
        .filter_by(session_id=session_id)
        .order_by(AIChatMessage.timestamp.asc())
        .all()
    )
    out = [{"role": m.role, "text": m.text, "timestamp": m.timestamp.isoformat()} for m in messages]
    return {"success": True, "messages": out, "title": session.title}


@router.delete("/ai/session/{session_id}")
def delete_ai_session(
    session_id: str,
    payload: Annotated[dict, Depends(get_token_payload)],
    db: Session = Depends(get_db),
):
    uid = int(payload.get("sub"))
    session = db.query(AIChatSession).filter_by(id=session_id, user_id=uid).first()
    if not session:
        return _fail("Session not found", 404)

    try:
        db.query(AIChatMessage).filter_by(session_id=session_id).delete()
        db.delete(session)
        db.commit()
        return {"success": True, "message": "Session deleted"}
    except Exception as e:
        print(f"Error deleting session: {e}")
        db.rollback()
        return _fail(f"Server error: {str(e)}", 500)


@router.get("/api/academic-insights")
def get_academic_insights(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    student_id = current_user.id

    marks = db.query(Mark).filter_by(student_id=student_id).all()
    marks_summary = [{"subject": m.subject, "score": m.marks_obtained, "max": m.max_marks} for m in marks]

    attendance_records = db.query(Attendance).filter_by(student_id=student_id).all()
    att_stats: dict = {}
    for r in attendance_records:
        if r.subject not in att_stats:
            att_stats[r.subject] = {"present": 0, "total": 0}
        att_stats[r.subject]["total"] += 1
        if r.status == "Present":
            att_stats[r.subject]["present"] += 1

    att_summary = []
    for subj, data in att_stats.items():
        pct = (data["present"] / data["total"]) * 100 if data["total"] > 0 else 0
        att_summary.append({"subject": subj, "percentage": round(pct, 1)})

    feedbacks = (
        db.query(Feedback)
        .filter_by(student_id=student_id)
        .order_by(Feedback.created_at.desc())
        .limit(3)
        .all()
    )
    fb_summary = [{"subject": f.subject, "text": f.text} for f in feedbacks]

    role_view = "Student" if current_user.role == "student" else "Parent"
    prompt = f"""
    Analyze the academic data for a student named {current_user.name}.
    Role View: {role_view} (Provide advice suitable for a {role_view}).

    Data:
    Marks: {json.dumps(marks_summary)}
    Attendance: {json.dumps(att_summary)}
    Recent Feedback: {json.dumps(fb_summary)}

    Task:
    1. Identify 'Attendance Risks' (Subjects < 75%).
    2. Identify 'Subject Priorities' (Low marks).
    3. Generate 'Improvement Suggestions' (Actionable steps).
    4. Write a 'Counselor Message': A warm, professional paragraph summarizing status and advice.
       - If Student view: Be encouraging, strategic.
       - If Parent view: Be reassuring, clear on what to monitor.

    Output strictly in valid JSON format:
    {{
        "attendance_risks": ["Subject A", ...],
        "priorities": ["Subject B", ...],
        "suggestions": ["Tip 1", "Tip 2", ...],
        "counselor_message": "..."
    }}
    """

    ai_response, _error = call_groq_api(prompt)
    if not ai_response:
        return {
            "success": True,
            "insights": {
                "attendance_risks": [],
                "priorities": [],
                "suggestions": ["Focus on consistent attendance.", "Review recent class notes."],
                "counselor_message": (
                    "AI services are currently unavailable, but please review your marks and attendance manually."
                ),
            },
        }

    try:
        clean_json = ai_response.replace("```json", "").replace("```", "").strip()
        parsed = json.loads(clean_json)
        return {"success": True, "insights": parsed}
    except Exception as e:
        print(f"AI Parse Error: {e}")
        return {
            "success": True,
            "insights": {
                "attendance_risks": [],
                "priorities": [],
                "suggestions": ["Please review your dashboard."],
                "counselor_message": ai_response,
            },
        }
